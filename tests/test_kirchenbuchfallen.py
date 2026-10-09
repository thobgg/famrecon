"""Fallen aus grossen echten Registern, an erfundenen Faellen festgehalten.

Ehename: "Anna, Kaspar Brenners Hausfrau" wird mit dem Namen des Mannes begraben; sie ist die Braut Anna Weber.
Vater = Ehemann: ein Begraebnis nennt denselben Mann als Vater und als Ehemann; daraus darf keine Ehe
zwischen Vater und Tochter werden. Namensweitergabe: ein Saeugling mit Alter in Tagen gehoert zur Taufe,
die zu diesem Alter passt, nicht zum aelteren Geschwister gleichen Namens. Schon begraben: wer ein
Sterbedatum hat, wird nicht ein zweites Mal begraben.
"""
import tempfile
import unittest
from pathlib import Path

import openpyxl

from famrecon import db, kern, lesen, normalform, verknuepfen, zuordnung


def projekt(ehen=(), taufen=(), tote=()):
    d = tempfile.mkdtemp()
    xlsx = Path(d) / "r.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Ehen"
    ws.append(["Trauungsdatum", "Bräutigam Name", "Bräutigam Vorname", "Braut Name", "Braut Vorname"])
    for r in ehen or [("1600-01-01", "Platzhalter", "Hans", "Platzhalter", "Eva")]:
        ws.append(list(r))
    ws = wb.create_sheet("Taufen")
    ws.append(["Taufdatum", "Name", "Vorname", "Vater", "Mutter"])
    for r in taufen or [("1600-01-01", "Platzhalter", "Hans", "Platzhalter, Hans", "Eva")]:
        ws.append(list(r))
    ws = wb.create_sheet("Tote")
    ws.append(["Sterbedatum", "Name", "Vorname", "Alter", "Stand", "Vater", "Ehepartner"])
    for r in tote:
        ws.append(list(r))
    wb.save(xlsx)
    toml = Path(d) / "z.toml"
    toml.write_text(zuordnung.toml_text(str(xlsx), zuordnung.vorschlagen(xlsx)), encoding="utf-8")
    con = db.oeffnen(Path(d) / "p.db")
    lesen.einlesen(con, lesen.zuordnung_laden(toml), [xlsx])
    kern.personen_bauen(con)
    verknuepfen.verknuepfen(con)
    return con


def ident_der_nennung(con, register, jahr, pfad):
    return con.execute("SELECT i.* FROM zuordnung z JOIN person p ON p.id=z.person JOIN eintrag e ON e.id=p.eintrag "
                       "JOIN identitaet i ON i.id=z.ident WHERE e.register=? AND e.jahr=? AND p.pfad=?",
                       (register, jahr, pfad)).fetchone()


class Kirchenbuchfallen(unittest.TestCase):
    def test_ehename_findet_die_braut(self):
        con = projekt(ehen=[("1600-05-10", "Brenner", "Kaspar", "Weber", "Anna")],
                      tote=[("1630-04-30", "Brenner", "Anna", None, "Ehefrau", None, "Brenner, Kaspar")])
        tote = ident_der_nennung(con, "tod", 1630, "verstorbener")
        braut = ident_der_nennung(con, "ehe", 1600, "braut")
        self.assertEqual(tote["id"], braut["id"])
        self.assertEqual(con.execute("SELECT COUNT(*) FROM familie f JOIN identitaet m ON m.id=f.mann "
                                     "WHERE m.name='Brenner'").fetchone()[0], 1)

    def test_ehename_ohne_trauung_bleibt_ohne_geburtsnamen(self):
        con = projekt(tote=[("1630-04-30", "Brenner", "Anna", None, "Witwe", None, "Brenner, Kaspar")])
        tote = ident_der_nennung(con, "tod", 1630, "verstorbener")
        self.assertIsNone(tote["name"])
        self.assertEqual(tote["ehename"], "Brenner")

    def test_vater_gleich_ehemann_keine_ehe_mit_dem_vater(self):
        con = projekt(taufen=[("1601-07-12", "Pfister", "Magdalena", "Pfister, Johann", "Anna")],
                      tote=[("1606-06-06", "Pfister", "Magdalena", None, "Witwe", "Pfister, Johann", "Pfister, Johann")])
        for f in con.execute("SELECT * FROM familie WHERE mann IS NOT NULL AND frau IS NOT NULL"):
            kinder = {r[0] for r in con.execute("SELECT ident FROM kind WHERE familie IN "
                                                "(SELECT id FROM familie WHERE mann=?)", (f["mann"],))}
            self.assertNotIn(f["frau"], kinder)
        # die Witwe ist nicht das fuenfjaehrige Kind von 1601
        self.assertNotEqual(ident_der_nennung(con, "tod", 1606, "verstorbener")["id"],
                            ident_der_nennung(con, "taufe", 1601, "kind")["id"])

    def test_namensweitergabe_saeugling_zum_passenden_geschwister(self):
        con = projekt(taufen=[("1620-09-29", "Lindner", "Christoph", "Lindner, Paul", "Maria"),
                              ("1622-06-10", "Lindner", "Christoph", "Lindner, Paul", "Maria")],
                      tote=[("1622-06-23", "Lindner", "Christoph", "2 Wochen", None, "Lindner, Paul", None)])
        self.assertEqual(ident_der_nennung(con, "tod", 1622, "verstorbener")["id"],
                         ident_der_nennung(con, "taufe", 1622, "kind")["id"])

    def test_schon_begraben(self):
        con = projekt(taufen=[("1596-03-01", "Holzapfel", "Anna", "Holzapfel, Johann", "Eva")],
                      tote=[("1607-09-16", "Holzapfel", "Anna", "11 Jahre", None, "Holzapfel, Johann", None),
                            ("1607-10-02", "Holzapfel", "Anna", None, None, None, None)])
        self.assertNotEqual(con.execute("SELECT z.ident FROM zuordnung z JOIN person p ON p.id=z.person JOIN eintrag e "
                                        "ON e.id=p.eintrag WHERE e.register='tod' AND e.monat=9").fetchone()[0],
                            con.execute("SELECT z.ident FROM zuordnung z JOIN person p ON p.id=z.person JOIN eintrag e "
                                        "ON e.id=p.eintrag WHERE e.register='tod' AND e.monat=10").fetchone()[0])

    def test_saeugling_nicht_zum_toten_vater(self):
        con = projekt(taufen=[("1598-10-09", "Rauch", "Katharina", "Rauch, Johann", "Margaretha")],
                      tote=[("1630-04-28", "Rauch", "Johann", "70 Jahre", None, None, None),
                            ("1632-02-07", "Rauch", "Margaretha", "8 Tage", None, "Rauch, Johann", None)])
        kind = ident_der_nennung(con, "tod", 1632, "verstorbener")
        vater = con.execute("SELECT mann FROM familie WHERE id=?", (kind["famc"],)).fetchone()[0]
        self.assertNotEqual(vater, ident_der_nennung(con, "tod", 1630, "verstorbener")["id"])

    def test_lautschluessel_wortende_und_lateinische_vornamen(self):
        self.assertEqual(normalform.koelner("Schmid"), normalform.koelner("Schmidt"))
        self.assertNotEqual(normalform.koelner("Schmid"), normalform.koelner("Schmitz"))
        for latein, deutsch in (("Georgius", "Georg"), ("Jacobus", "Jakob"), ("Petrus", "Peter"), ("Jeorgius", "Georg")):
            self.assertEqual(normalform.vorname_kanon(latein), normalform.vorname_kanon(deutsch))
        self.assertEqual(normalform.vorname_kanon("Nicolaus"), normalform.vorname_kanon("Nikolaus"))
        self.assertEqual(normalform.stand_art("Witwe"), "ehe")
        self.assertEqual(normalform.stand_art("Säugling"), "kind")


if __name__ == "__main__":
    unittest.main()
