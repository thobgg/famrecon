"""Urteilstabelle hin und zurueck: Pruefliste als xlsx, Urteil eintragen, als Entscheidung einlesen, Lauf folgt ihm."""
import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from famrecon import db, kern, lesen, urteile, verknuepfen

B = Path(__file__).resolve().parent.parent / "beispiel"


class Urteile(unittest.TestCase):
    def setUp(self):
        self.con = db.oeffnen(":memory:")
        lesen.einlesen(self.con, lesen.zuordnung_laden(B / "kirchenbuchstil.toml"), [B / "kirchenbuchstil.xlsx"])
        kern.personen_bauen(self.con)
        verknuepfen.verknuepfen(self.con)

    def test_hin_und_zurueck(self):
        with tempfile.TemporaryDirectory() as d:
            pfad = Path(d) / "pruefliste.xlsx"
            n = urteile.tabelle_schreiben(self.con, pfad, alle=True)
            self.assertGreater(n, 0)
            wb = openpyxl.load_workbook(pfad); ws = wb.active
            kopf = [c.value for c in ws[1]]
            self.assertEqual(kopf[:2], ["FALL", "SCHLUESSEL"]); self.assertIn("URTEIL", kopf)
            i_u, i_k, i_s, i_w = kopf.index("URTEIL") + 1, kopf.index("KANDIDATEN") + 1, kopf.index("SCHLUESSEL") + 1, kopf.index("WAHL") + 1
            # Zeile 2: Urteil "neu"; Zeile 3: den ersten Kandidaten bestaetigen; Zeile 4: offen lassen
            ws.cell(row=2, column=i_u, value="neu")
            kand = ws.cell(row=3, column=i_k).value
            nummer = kand.split("]")[0].strip("[")
            ws.cell(row=3, column=i_u, value=nummer)
            ws.cell(row=4, column=i_u, value="offen")
            wb.save(pfad)
            st = urteile.tabelle_lesen(self.con, pfad)
            self.assertEqual((st["neu"], st["gleich"]), (1, 1), st)
            self.assertEqual(st["fehler"], [])
            self.assertEqual(self.con.execute("SELECT COUNT(*) FROM entscheidung").fetchone()[0], 2)
            schl2 = ws.cell(row=2, column=i_s).value
            self.assertEqual(self.con.execute("SELECT art FROM entscheidung WHERE schluessel=?", (schl2,)).fetchone()[0], "neu")
            # der naechste Lauf folgt den Urteilen
            verknuepfen.verknuepfen(self.con)
            r = self.con.execute("SELECT z.stufe, z.grund FROM zuordnung z JOIN person p ON p.id=z.person JOIN eintrag e ON e.id=p.eintrag "
                                 "JOIN quelle q ON q.id=e.quelle WHERE q.datei||'|'||q.blatt||'|'||e.zeile||'|'||p.pfad=?", (schl2,)).fetchone()
            self.assertIn("von Hand", r["grund"])
            # zuruecknehmen
            ws.cell(row=2, column=i_u, value="zurueck"); ws.cell(row=3, column=i_u, value=""); wb.save(pfad)
            st = urteile.tabelle_lesen(self.con, pfad)
            self.assertEqual(st["zurueck"], 1)
            self.assertEqual(self.con.execute("SELECT COUNT(*) FROM entscheidung").fetchone()[0], 1)


if __name__ == "__main__":
    unittest.main()
