"""Stichprobe in Kopien der Tabellen: Inhalt unveraendert, Familien farbig, Auswertung zaehlt richtig."""
import shutil
import tempfile
import unittest
from pathlib import Path

import openpyxl

from famrecon import db, kern, lesen, stichprobe, verknuepfen

B = Path(__file__).resolve().parent.parent / "beispiel"


class Stichprobe(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = Path(tempfile.mkdtemp())
        shutil.copy(B / "kirchenbuchstil.xlsx", cls.d / "kirchenbuchstil.xlsx")
        cls.con = db.oeffnen(cls.d / "p.db")
        lesen.einlesen(cls.con, lesen.zuordnung_laden(B / "kirchenbuchstil.toml"), [cls.d / "kirchenbuchstil.xlsx"])
        kern.personen_bauen(cls.con)
        verknuepfen.verknuepfen(cls.con)
        cls.st = stichprobe.bauen(cls.con, cls.d, cls.d / "s", n_zufall=6, n_schwer=0, saat=3)

    def test_inhalt_unveraendert_und_spalten(self):
        alt = openpyxl.load_workbook(self.d / "kirchenbuchstil.xlsx")
        neu = openpyxl.load_workbook(self.d / "s" / "kirchenbuchstil.xlsx")
        for wa in alt.worksheets:
            wn = neu[wa.title]
            for r in range(1, wa.max_row + 1):
                for c in range(1, wa.max_column + 1):
                    self.assertEqual(wa.cell(r, c).value, wn.cell(r, c).value)
            self.assertEqual([wn.cell(1, wa.max_column + k).value for k in range(1, 6)], stichprobe.NEUE_SPALTEN)
            for r in range(2, wa.max_row + 1):
                if any(wa.cell(r, c).value is not None for c in range(1, wa.max_column + 1)):
                    self.assertTrue(wn.cell(r, wa.max_column + 2).value, (wa.title, r))     # Person-ID in jeder Datenzeile

    def test_farben_je_familie(self):
        neu = openpyxl.load_workbook(self.d / "s" / "kirchenbuchstil.xlsx")
        farbe_von = {}
        for ws in neu.worksheets:
            sp = [c.value for c in ws[1]].index("Stichprobe") + 1
            for r in range(2, ws.max_row + 1):
                f = ws.cell(r, sp).value
                if f:
                    farbe = ws.cell(r, 1).fill.fgColor.rgb
                    farbe_von.setdefault(f.split(", ")[0], set()).add(farbe)
        self.assertEqual(len(farbe_von), self.st["familien"])
        self.assertTrue(all(len(v) == 1 for v in farbe_von.values()))                  # eine Farbe je Familie, ueber alle Blaetter
        self.assertEqual(len({next(iter(v)) for v in farbe_von.values()}), len(farbe_von))   # verschiedene Familien, verschiedene Farben

    def test_auswertung(self):
        aus = self.d / "aus"
        shutil.copytree(self.d / "s", aus)
        wb = openpyxl.load_workbook(aus / "kirchenbuchstil.xlsx")
        erste_falsch = None
        for ws in wb.worksheets:
            k = {c.value: c.column for c in ws[1]}
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, k["Stichprobe"]).value:
                    u = "richtig"
                    if erste_falsch is None:
                        u, erste_falsch = "falsch", ws.cell(r, k["Stichprobe"]).value.split(", ")[0]
                    ws.cell(r, k["Gehört dazu?"], u)
        wb.save(aus / "kirchenbuchstil.xlsx")
        ub = openpyxl.load_workbook(aus / "Stichprobe Übersicht.xlsx")
        ub["Familien"].cell(3, 7, "ja"); ub["Familien"].cell(3, 8, "Taufen, Zeile 4")
        ub.save(aus / "Stichprobe Übersicht.xlsx")
        r = stichprobe.auswerten(aus)
        z = r["zufall"]
        self.assertEqual(z["falsch"], 1)
        self.assertEqual(z["richtig"] + z["falsch"], self.st["zeilen"])
        self.assertIn(erste_falsch, r["fremde"])
        self.assertEqual(len(r["fehlt"]), 1)
        self.assertGreaterEqual(z["familien_fehlerhaft"], 1)


if __name__ == "__main__":
    unittest.main()
