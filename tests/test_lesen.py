"""Einlesen der Kirchenbuchstil- Beispieldatei: Zaehler, Sammelfelder, Leerwoerter."""
import unittest
from pathlib import Path

from famrecon import db, lesen

HIER = Path(__file__).resolve().parent.parent
XLSX = HIER / "beispiel" / "kirchenbuchstil.xlsx"
TOML = HIER / "beispiel" / "kirchenbuchstil.toml"


class Einlesen(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.con = db.oeffnen(":memory:")
        cls.zaehler = lesen.einlesen(cls.con, lesen.zuordnung_laden(TOML), [XLSX])

    def feld(self, register, feldname, **wo):
        bed, par = [], []
        for k, v in wo.items():
            bed.append("e.id IN (SELECT eintrag FROM feld WHERE name=? AND wert=?)")
            par += [k, v]
        sql = ("SELECT f.wert, f.roh FROM eintrag e JOIN feld f ON f.eintrag=e.id "
               f"WHERE e.register=? AND f.name=? AND {' AND '.join(bed)}")
        return self.con.execute(sql, [register, feldname] + par).fetchone()

    def test_zaehler(self):
        self.assertEqual(self.zaehler, {"taufe": 7, "ehe": 14, "tod": 14})

    def test_jahre(self):
        r = self.con.execute("SELECT register, MIN(jahr), MAX(jahr) FROM eintrag "
                             "GROUP BY register ORDER BY register").fetchall()
        self.assertEqual([tuple(x) for x in r],
                         [("ehe", 1732, 1804), ("taufe", 1727, 1815), ("tod", 1722, 1815)])
        self.assertEqual(self.con.execute("SELECT COUNT(*) FROM eintrag WHERE jahr IS NULL").fetchone()[0], 0)

    def test_paten_gesammelt(self):
        wert, roh = self.feld("taufe", "paten", kind_name="Haag")
        self.assertEqual(len(wert.split(lesen.TRENNER)), 5)
        self.assertTrue(wert.startswith("Weiß, Christoph Friedrich"))

    def test_leerwort_bleibt_roh(self):
        wert, roh = self.feld("ehe", "braut_vater", braeutigam_name="Feldmann")
        self.assertIsNone(wert)
        self.assertEqual(roh, "k.A.")
        wert, roh = self.feld("taufe", "paten", kind_name="Ebel")
        self.assertIsNone(wert)
        self.assertEqual(roh, "—")

    def test_zwei_spalten_ein_feld(self):
        wert, _ = self.feld("ehe", "bemerkung", braeutigam_name="Obermüller (W)")
        self.assertEqual(len(wert.split(lesen.TRENNER)), 1)      # nur Bemerkung 2 gefuellt
        wert, _ = self.feld("ehe", "bemerkung", braeutigam_name="Rosenfeld von")
        self.assertTrue(wert.startswith("nach diesem Eintrag"))

    def test_zeilenumbruch_im_kb(self):
        wert, _ = self.feld("taufe", "kb", kind_name="Arst")
        self.assertEqual(wert, "Taufen 1738-1758")

    def test_alter_null_als_text(self):
        wert, _ = self.feld("tod", "verstorbener_alter", verstorbener_name="Ebel")
        self.assertEqual(wert, "0")


if __name__ == "__main__":
    unittest.main()


class CSV(unittest.TestCase):
    """CSV mit Semikolon und BOM, wie deutsches Excel es schreibt: ein Blatt je Datei."""

    def test_falkenrath_csv(self):
        b = lesen.blaetter(HIER / "beispiel" / "falkenrath-taufen.csv")
        self.assertEqual(len(b), 1)
        blatt, zeilen, kopf = b[0]
        self.assertEqual((blatt, zeilen), ("falkenrath-taufen", 492))
        self.assertEqual(kopf[0], ("A", "Geburtsdatum"))
        erste = next(lesen.zeilen_lesen(HIER / "beispiel" / "falkenrath-taufen.csv", "falkenrath-taufen"))
        self.assertEqual(erste[0], 2)
        self.assertRegex(str(erste[1][0]), r"^\d{4}-\d{2}-\d{2}$")


class ZweiBlaetterEinRegister(unittest.TestCase):
    """Zwei Taufblaetter (zwei Baende) in einer Mappe: die Zuordnung muss eindeutige Abschnitte schreiben
    (taufe, taufe_2) und das Einlesen beide als Register taufe lesen."""

    def test_zwei_taufblaetter(self):
        import openpyxl, tempfile
        from famrecon import zuordnung as zu
        quelle = openpyxl.load_workbook(XLSX)
        with tempfile.TemporaryDirectory() as d:
            ziel = Path(d) / "zwei.xlsx"
            wb = openpyxl.Workbook(); wb.remove(wb.active)
            for titel, von in (("Taufen 1700", "Taufen"), ("Taufen 1750", "Taufen"), ("Ehen", "Ehen"), ("Tote", "Tote")):
                ws = wb.create_sheet(titel)
                for row in quelle[von].iter_rows(values_only=True):
                    ws.append(list(row))
            wb.save(ziel)
            text = zu.toml_text(str(ziel), zu.vorschlagen(ziel))
            self.assertIn("[register.taufe_2]", text)
            self.assertIn('register = "taufe"', text)
            toml = Path(d) / "z.toml"; toml.write_text(text, encoding="utf-8")
            z = lesen.zuordnung_laden(toml)                                # kein 'Cannot declare twice' mehr
            self.assertEqual(sorted(r["register"] for r in z["register"].values()), ["ehe", "taufe", "taufe", "tod"])
            con = db.oeffnen(":memory:")
            n = lesen.einlesen(con, z, [ziel])
            self.assertEqual(n["taufe"], 14)                               # 7 Taufen je Blatt, beide gelesen
            self.assertEqual(con.execute("SELECT COUNT(*) FROM quelle WHERE register='taufe'").fetchone()[0], 2)


class ZusatzUndKreuz(unittest.TestCase):
    """Erfundene Taufliste: Umschrift und Spalte ohne Feld stehen im Text der Fundstelle, X = "" bleibt weg, m/w-Kreuze ergeben das Geschlecht."""
    CSV = ("Taufdatum;Vorname;Name;m;w;Vater Vorname;Vater Name;Mutter Vorname;Mutter Name;Hausnummer;Notiz;Umschrift\n"
           "1781-03-04;Anna;Kessler;;x;Jakob;Kessler;Maria;Bühler;12;nur intern;Den 4. Mart. ist getaufft worden Anna\n"
           "1782-05-06;Kim;Kessler;x;;Jakob;Kessler;Maria;Bühler;;;\n"
           "1783-07-08;Kim;Kessler;x;x;Jakob;Kessler;Maria;Bühler;;;\n")
    TOML = ('[allgemein]\nleer = []\n\n[register.taufe]\nblatt = "taufen"\nleitdatum = ["tauf_datum"]\n\n'
            '[register.taufe.spalten]\nA = "tauf_datum"\nB = "kind_vorname"\nC = "kind_name"\n'
            'D = "kind_geschlecht=m"\nE = "kind_geschlecht=w"\nF = "vater_vorname"\nG = "vater_name"\n'
            'H = "mutter_vorname"\nI = "mutter_name"\nK = ""\nL = "wortlaut"\n')

    @classmethod
    def setUpClass(cls):
        import tempfile
        from famrecon import gedcom, kern, verknuepfen
        cls.tmp = tempfile.TemporaryDirectory()
        d = Path(cls.tmp.name)
        (d / "taufen.csv").write_text(cls.CSV, encoding="utf-8")
        (d / "z.toml").write_text(cls.TOML, encoding="utf-8")
        cls.con = db.oeffnen(":memory:")
        lesen.einlesen(cls.con, lesen.zuordnung_laden(d / "z.toml"), [d / "taufen.csv"])
        kern.personen_bauen(cls.con)
        verknuepfen.verknuepfen(cls.con)
        gedcom.schreiben(cls.con, d / "t.ged")
        cls.ged = (d / "t.ged").read_text(encoding="utf-8")

    @classmethod
    def tearDownClass(cls):
        cls.con.close()
        cls.tmp.cleanup()

    def felder(self, zeile):
        return {r["name"]: r["wert"] for r in self.con.execute(
            "SELECT f.name, f.wert FROM feld f JOIN eintrag e ON e.id=f.eintrag WHERE e.zeile=?", (zeile,))}

    def test_zusatz_und_weglassen(self):
        self.assertNotIn("nur intern", self.ged)                                               # K = "": weggelassen
        self.assertIn("3 DATA\n4 TEXT Den 4. Mart. ist getaufft worden Anna\n5 CONT Hausnummer: 12\n", self.ged)
        self.assertNotIn("TEXT Taufdatum", self.ged)                       # zugeordnete Spalten stehen nicht doppelt
        self.assertEqual(self.ged.count("3 DATA"), 1)                      # die anderen Taufen haben nichts ohne Feld
        zellen = self.con.execute("SELECT COUNT(*) FROM zelle z JOIN eintrag e ON e.id=z.eintrag WHERE e.zeile=2").fetchone()[0]
        self.assertEqual(zellen, 10)                                       # alle nicht leeren Zellen der Zeile, ohne K

    def test_kreuzspalten(self):
        self.assertEqual(self.felder(2)["kind_geschlecht"], "w")
        self.assertEqual(self.felder(3)["kind_geschlecht"], "m")
        self.assertEqual(self.felder(4)["kind_geschlecht"], "m | w")      # beides angekreuzt: kein eindeutiger Wert
        g = {r["zeile"]: r["geschlecht"] for r in self.con.execute(
            "SELECT e.zeile, p.geschlecht FROM person p JOIN eintrag e ON e.id=p.eintrag WHERE p.pfad='kind'")}
        self.assertEqual((g[2], g[3]), ("F", "M"))

    def test_keine_warnung_fuer_feste_werte(self):
        self.assertEqual(lesen.unbekannte_felder(lesen.zuordnung_laden(Path(self.tmp.name) / "z.toml")), [])
