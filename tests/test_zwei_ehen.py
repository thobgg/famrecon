"""Ein Witwer mit zwei Frauen gleichen Vornamens; in der Taufe steht die Mutter nur als "N., Catharina".

Ist der Tod der ersten Frau bekannt, entscheidet er. Sonst passen beide Ehen gleich gut:
die juengere Ehe wird genommen und der Fall geht in die Pruefliste. Ein Urteil dazu gilt beim naechsten Lauf.
"""
import tempfile
import unittest
from pathlib import Path

import openpyxl

from famrecon import db, kern, lesen, verknuepfen, zuordnung


def projekt(ehen, taufen, tote=()):
    d = tempfile.mkdtemp()
    xlsx = Path(d) / "r.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Ehen"
    ws.append(["Trauungsdatum", "Bräutigam Name", "Bräutigam Vorname", "Braut Name", "Braut Vorname"])
    for r in ehen:
        ws.append(list(r))
    ws = wb.create_sheet("Taufen")
    ws.append(["Taufdatum", "Name", "Vorname", "Vater", "Mutter"])
    for r in taufen:
        ws.append(list(r))
    ws = wb.create_sheet("Tote")
    ws.append(["Sterbedatum", "Name", "Vorname", "Ehepartner"])
    for r in tote or [("1700-01-01", "Platzhalter", "Hans", None)]:
        ws.append(list(r))
    wb.save(xlsx)
    toml = Path(d) / "z.toml"
    toml.write_text(zuordnung.toml_text(str(xlsx), zuordnung.vorschlagen(xlsx)), encoding="utf-8")
    con = db.oeffnen(Path(d) / "p.db")
    lesen.einlesen(con, lesen.zuordnung_laden(toml), [xlsx])
    kern.personen_bauen(con)
    verknuepfen.verknuepfen(con)
    return con


EHEN = [("1758-05-10", "Kiefer", "Friedrich", "Burckhardt", "Catharina"),
        ("1767-02-03", "Kiefer", "Friedrich", "Haas", "Catharina")]
TAUFE_1770 = ("1770-03-01", "Kiefer", "Johann", "Kiefer, Friedrich", "N., Catharina")


def familie_des_kindes(con, vorname):
    return con.execute("SELECT f.* FROM kind k JOIN identitaet i ON i.id=k.ident JOIN familie f ON f.id=k.familie "
                       "WHERE i.vorname=?", (vorname,)).fetchone()


def zuordnung_mutter_1770(con):
    return con.execute("SELECT z.*, p.id pid FROM zuordnung z JOIN person p ON p.id=z.person JOIN eintrag e ON e.id=p.eintrag "
                       "WHERE e.register='taufe' AND e.jahr=1770 AND p.pfad='mutter'").fetchone()


class ZweiEhen(unittest.TestCase):
    def test_tod_der_ersten_frau_entscheidet(self):
        con = projekt(EHEN, [TAUFE_1770], [("1766-08-20", "Kiefer geb. Burckhardt", "Catharina", "Kiefer, Friedrich")])
        f = familie_des_kindes(con, "Johann")
        self.assertEqual(f["tr_jahr"], 1767)
        self.assertNotEqual(zuordnung_mutter_1770(con)["stufe"], "unsicher")

    def test_gleichstand_geht_in_die_pruefliste(self):
        con = projekt(EHEN, [TAUFE_1770])
        f = familie_des_kindes(con, "Johann")
        self.assertIn(f["tr_jahr"], (1758, 1767))
        self.assertEqual(f["trauung_eintrag"] is not None, True)           # keine dritte, erschlossene Familie
        z = zuordnung_mutter_1770(con)
        self.assertEqual(z["stufe"], "unsicher")
        self.assertIn("gleich gut", z["grund"])
        self.assertTrue(z["alternativen"] and z["alternativen"] != "[]")
        self.assertEqual(con.execute("SELECT COUNT(*) FROM familie WHERE mann IS NOT NULL").fetchone()[0], 2)

    def test_urteil_gilt_beim_naechsten_lauf(self):
        con = projekt(EHEN, [TAUFE_1770])
        z = zuordnung_mutter_1770(con)
        andere = 1758 if familie_des_kindes(con, "Johann")["tr_jahr"] == 1767 else 1767
        braut = con.execute("SELECT p.id FROM person p JOIN eintrag e ON e.id=p.eintrag WHERE e.register='ehe' AND e.jahr=? "
                            "AND p.pfad='braut'", (andere,)).fetchone()["id"]
        verknuepfen.entscheiden_von_hand(con, z["pid"], "gleich", braut)
        verknuepfen.verknuepfen(con)
        self.assertEqual(familie_des_kindes(con, "Johann")["tr_jahr"], andere)
        self.assertEqual(con.execute("SELECT COUNT(*) FROM familie WHERE mann IS NOT NULL").fetchone()[0], 2)


if __name__ == "__main__":
    unittest.main()
