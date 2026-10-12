"""Stichprobe zur Pruefung der Zuordnung, in den Tabellen des Erfassers selbst ("wie mit Leuchtmarker").

    famrecon stichprobe daten/x/projekt.db --quelle daten/x -o daten/x/stichprobe
    famrecon stichprobe-auswerten daten/x/stichprobe            (nach Rueckgabe der ausgefuellten Dateien)

Erzeugt Kopien der Registertabellen, inhaltlich unveraendert (Zellen, Kommentare, eingefuegte Bilder), mit fuenf
Spalten am rechten Rand: Stichprobe (Familiennummer), Person-ID, Familien-ID, Gehoert dazu? (richtig / falsch /
unklar), Kommentar. Person- und Familien-ID stehen in jeder Zeile; die Zeilen der gezogenen Familien sind in einer
Farbe je Familie hinterlegt, ueber alle Dateien hinweg. Eine Uebersicht nennt je Familie Farbe, Personen und die
Zeilen und fragt, ob eine Zeile fehlt.

Eine Zeile gehoert zu einer Familie: die Trauung des Paares, die Taufen ihrer Kinder, die Begraebnisse von Mann,
Frau und Kindern. Gezogen wird nur unter Familien mit mindestens zwei Zeilen: nur dort wurde etwas verbunden,
eine Familie aus einer einzigen Zeile kann keine falsche Verbindung enthalten. Keine Zeile gehoert zu zwei
gezogenen Familien, damit jede Zeile genau eine Farbe hat. Dazu je Familie die naheliegenden Zeilen, die das Programm NICHT zugeordnet hat (Taufen mit gleichnamigem Vater im
Zeitraum der Familie, Begraebnisse gleichnamiger Mitglieder ohne eigenes Begraebnis), in einem hellen Ton der
Familienfarbe und mit derselben Frage: so wird "fehlt eine Zeile?" vom Suchen zum Ankreuzen, und die Vollstaendigkeit
laesst sich messen. Ein Filter auf der Spalte Stichprobe zeigt nur die Stichprobe.

Die Antworten wirken direkt (stichprobe-auswerten --db ... --uebernehmen): Kandidat "richtig" -> Entscheidung
"gleich" (Vater/Mutter einer Taufe = Mann/Frau der Familie; Verstorbener = das Familienmitglied); Begraebnis "falsch"
-> Entscheidung "neu" (eigene Person). Falsche Taufen und Trauungen bleiben fuer die Pruefliste, weil offen ist,
welche Verbindung falsch war.

Gezogen werden Familien zufaellig (die Grundlage der Fehlerquote) und bewusst schwierige Faelle
(haeufige Namen, Witwer mit zweiter Ehe, Taufen mit einem Paten gleichen Namens wie der Vater), die getrennt
ausgewertet werden. Die Abschrift selbst wird nicht geprueft.
"""
import io
import json
import math
import random
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.drawing.image import Image as XlBild
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor, TwoCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

NEUE_SPALTEN = ["Stichprobe", "Person-ID", "Familien-ID", "Gehört dazu?", "Kommentar"]
REG = {"taufe": "Taufe", "ehe": "Trauung", "tod": "Begräbnis"}
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
      "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def hell(rgb, anteil=0.6):
    """Farbe zu anteil mit Weiss mischen: der helle Ton fuer Kandidatenzeilen."""
    r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    return "%02X%02X%02X" % tuple(int(x + (255 - x) * anteil) for x in (r, g, b))


def farben(n):
    """n gut unterscheidbare helle Farben (Goldener Winkel ueber den Farbkreis, zwei Helligkeiten)."""
    import colorsys
    out = []
    for k in range(n):
        hue = (k * 0.618033988749895) % 1.0
        l, s = (0.80, 0.75) if k % 2 == 0 else (0.88, 0.70)
        r, g, b = colorsys.hls_to_rgb(hue, l, s)
        out.append("%02X%02X%02X" % (int(r * 255), int(g * 255), int(b * 255)))
    return out


# ------------------------------------------------------------- Daten
class Bestand:
    def __init__(self, con):
        self.quellen = {q["id"]: dict(q) for q in con.execute("SELECT * FROM quelle")}
        self.eintraege = {e["id"]: dict(e) for e in con.execute("SELECT * FROM eintrag")}
        self.idents = {i["id"]: dict(i) for i in con.execute("SELECT * FROM identitaet")}
        self.fams = {f["id"]: dict(f) for f in con.execute("SELECT * FROM familie")}
        self.kinder = defaultdict(list)
        for k in con.execute("SELECT familie, ident FROM kind"):
            self.kinder[k["familie"]].append(k["ident"])
        self.als_partner = defaultdict(list)
        for f in self.fams.values():
            for p in (f["mann"], f["frau"]):
                if p:
                    self.als_partner[p].append(f["id"])
        self.rolle = defaultdict(dict)                       # eintrag -> pfad -> ident
        for r in con.execute("SELECT p.eintrag, p.pfad, z.ident FROM person p JOIN zuordnung z ON z.person=p.id"):
            self.rolle[r["eintrag"]][r["pfad"]] = r["ident"]
        self.nennung = defaultdict(dict)                     # eintrag -> pfad -> person-Zeile
        for r in con.execute("SELECT id, eintrag, pfad, name, vorname, name_schl, vorname_kanon FROM person"):
            self.nennung[r["eintrag"]][r["pfad"]] = dict(r)
        self.erste_nennung = {}                              # ident -> eine person.id, die ihm zugeordnet ist
        for r in con.execute("SELECT ident, MIN(person) p FROM zuordnung GROUP BY ident"):
            self.erste_nennung[r["ident"]] = r["p"]
        self.paten = {r["eintrag"]: r["wert"] for r in con.execute("SELECT eintrag, wert FROM feld WHERE name='paten' AND wert IS NOT NULL")}
        self.trauung = {f["trauung_eintrag"]: f["id"] for f in self.fams.values() if f["trauung_eintrag"]}

    def ort(self, eid):
        e = self.eintraege[eid]
        q = self.quellen[e["quelle"]]
        return q["datei"], q["blatt"], e["zeile"]

    def ids(self, eid):
        """(Person-ID, Familien-ID) einer Tabellenzeile als Text."""
        e, r = self.eintraege[eid], self.rolle[eid]
        if e["register"] == "taufe":
            k = r.get("kind")
            fam = self.idents[k]["famc"] if k else None
            return (f"P{k}" if k else ""), (f"F{fam}" if fam else "")
        if e["register"] == "ehe":
            pers = " / ".join(f"P{r[x]}" for x in ("braeutigam", "braut") if r.get(x))
            fam = self.trauung.get(eid)
            return pers, (f"F{fam}" if fam else "")
        v = r.get("verstorbener")
        if not v:
            return "", ""
        fams = self.als_partner.get(v) or ([self.idents[v]["famc"]] if self.idents[v]["famc"] else [])
        return f"P{v}", ", ".join(f"F{f}" for f in fams)

    def zeilen_der_familie(self, fid):
        """Eintraege, die zu Familie fid gehoeren: Trauung, Taufen der Kinder, Begraebnisse von Mann, Frau, Kindern."""
        f = self.fams[fid]
        kinder = set(self.kinder.get(fid, []))
        leute = kinder | {x for x in (f["mann"], f["frau"]) if x}
        out = []
        for eid, r in self.rolle.items():
            reg = self.eintraege[eid]["register"]
            if (reg == "taufe" and r.get("kind") in kinder) or (reg == "ehe" and eid == f["trauung_eintrag"]) \
                    or (reg == "tod" and r.get("verstorbener") in leute):
                out.append(eid)
        return sorted(out, key=lambda e: (self.eintraege[e]["jahr"] or 0, e))

    def name(self, i):
        if not i:
            return "—"
        x = self.idents[i]
        return f"{x['name'] or ('NN' if x['unbekannt'] else '—')}, {x['vorname'] or '—'}"


def _passt(nennung, ident, B):
    """Name aehnlich (auch Ehe- oder Geburtsname) und Vorname nicht widersprechend."""
    from . import normalform as nf
    if not nennung or not ident:
        return False
    i = B.idents[ident]
    def nah(a, b):                                           # Lautschluessel UND Schreibung nah (Lemke ist nicht Lang)
        return bool(a and b) and nf.namen_aehnlich(a, b) and nf.levenshtein(a.lower(), b.lower()) <= max(2, len(a) // 3)
    if not any(nah(n, nennung.get("name")) for n in (i["name"], i["ehename"], i["geburtsname"]) if n):
        return False
    vp = nf.vornamen_punkte(nennung.get("vorname_kanon"), i["vorname_kanon"])
    return vp is None or vp > 0


def kandidaten(B, fid, belegt, je_familie=4):
    """Naheliegende, NICHT zugeordnete Zeilen einer Familie: [(eintrag, art, ziel_ident)].
    Taufen mit gleichnamigem Vater (und nicht widersprechender Mutter) im Zeitraum der Familie;
    Begraebnisse gleichnamiger Mitglieder, die noch kein Begraebnis haben."""
    f = B.fams[fid]
    eigene = set(B.zeilen_der_familie(fid))
    kinder = B.kinder.get(fid, [])
    jahre = [B.idents[k]["geb_jahr"] for k in kinder if B.idents[k]["geb_jahr"]] + ([f["tr_jahr"]] if f["tr_jahr"] else [])
    if not jahre:
        return []
    von, bis = min(jahre) - 3, max(jahre) + 12
    begraben = {B.rolle[e].get("verstorbener") for e in eigene if B.eintraege[e]["register"] == "tod"}
    out = []
    for eid, e in B.eintraege.items():
        if eid in eigene or eid in belegt or not e["jahr"]:
            continue
        n = B.nennung.get(eid, {})
        if e["register"] == "taufe" and von <= e["jahr"] <= bis and _passt(n.get("vater"), f["mann"], B):
            if f["frau"] and n.get("mutter") and n["mutter"].get("vorname_kanon") and not _passt(n["mutter"], f["frau"], B) \
                    and B.idents[f["frau"]]["vorname_kanon"]:
                from . import normalform as nf
                if not nf.vornamen_punkte(n["mutter"]["vorname_kanon"], B.idents[f["frau"]]["vorname_kanon"]):
                    continue
            out.append((abs(e["jahr"] - (von + bis) / 2), eid, "Taufe eines Kindes", f["mann"]))
        elif e["register"] == "tod":
            for m in [x for x in (f["mann"], f["frau"]) if x] + kinder:
                if m in begraben:
                    continue
                gj = B.idents[m]["geb_jahr"]
                if gj and not (gj <= e["jahr"] <= gj + 100):
                    continue
                if _passt(n.get("verstorbener"), m, B):
                    out.append((abs(e["jahr"] - (von + bis) / 2), eid, "Begräbnis eines Mitglieds", m))
                    break
    out.sort()
    return [(eid, art, ziel) for _, eid, art, ziel in out[:je_familie]]


def auswahl(B, n_zufall=20, n_schwer=10, saat=1):
    """-> [(familie, art)], art = zufall | haeufiger Name | Witwer, zweite Ehe | Pate gleichen Namens."""
    rnd = random.Random(saat)
    brauchbar = [f for f in B.fams.values() if f["mann"] and (f["trauung_eintrag"] or B.kinder.get(f["id"]))]
    gewaehlt, schon, belegt = [], set(), set()

    def frei(fid):
        """Familie hat eigene Zeilen, die noch keiner gezogenen Familie gehoeren (eine Farbe je Zeile)."""
        z = set(B.zeilen_der_familie(fid))
        return len(z) >= 2 and not (z & belegt)          # nur wo verbunden wurde, gibt es etwas zu pruefen

    def dazu(fid, art):
        gewaehlt.append((fid, art)); schon.add(fid); belegt.update(B.zeilen_der_familie(fid))

    reihe = brauchbar[:]
    rnd.shuffle(reihe)
    for f in reihe:
        if sum(1 for _, a in gewaehlt if a == "zufall") >= n_zufall:
            break
        if frei(f["id"]):
            dazu(f["id"], "zufall")

    def nimm(kand, anzahl, art):
        kand = [f for f in kand if f not in schon]
        rnd.shuffle(kand)
        for f in kand:
            if anzahl <= 0:
                break
            if frei(f):
                dazu(f, art); anzahl -= 1

    maenner = Counter(B.idents[f["mann"]]["name"] for f in brauchbar)
    vornamen = Counter((B.idents[f["mann"]]["vorname_kanon"] or "").split(" ")[0] for f in brauchbar)
    top_n = {n for n, _ in maenner.most_common(15)}
    top_v = {v for v, _ in vornamen.most_common(10) if v}
    haeufig = [f["id"] for f in brauchbar if B.idents[f["mann"]]["name"] in top_n
               and (B.idents[f["mann"]]["vorname_kanon"] or "").split(" ")[0] in top_v]
    witwer = defaultdict(list)
    for f in brauchbar:
        if f["frau"]:
            witwer[f["mann"]].append(f["id"])
    witwer = [fs for fs in witwer.values() if len(fs) >= 2 and any(B.kinder.get(x) for x in fs)]
    paten = []
    for f in brauchbar:
        m = B.idents[f["mann"]]
        if not m["name"] or not m["vorname"]:
            continue
        rufname = m["vorname"].split()[0].lower()
        for k in B.kinder.get(f["id"], []):
            for eid, r in B.rolle.items():
                if r.get("kind") == k and eid in B.paten:
                    p = B.paten[eid].lower()
                    if m["name"].lower() in p and rufname in p:
                        paten.append(f["id"])
    paten = list(dict.fromkeys(paten))

    rest = n_schwer
    n_z = sum(1 for _, a in gewaehlt if a == "zufall")
    nimm(haeufig, min(4, rest), "häufiger Name"); rest = n_schwer - (len(gewaehlt) - n_z)
    rnd.shuffle(witwer)
    for fs in witwer:
        if rest < 2:
            break
        zwei = sorted(fs, key=lambda x: (B.fams[x]["tr_jahr"] or 0))[:2]
        if not any(x in schon for x in zwei) and frei(zwei[0]) and frei(zwei[1]) \
                and not (set(B.zeilen_der_familie(zwei[0])) & set(B.zeilen_der_familie(zwei[1]))):
            for x in zwei:
                dazu(x, "Witwer, zweite Ehe")
            rest -= 2
        if sum(1 for _, a in gewaehlt if a.startswith("Witwer")) >= 4:
            break
    nimm(paten, max(0, rest), "Pate gleichen Namens")
    grund = sum(1 for f in brauchbar if len(B.zeilen_der_familie(f["id"])) >= 2)
    return gewaehlt, grund


# ------------------------------------------------------------- Bilder erhalten
def _bilder_lesen(pfad):
    """Eingefuegte Bilder je Blatt aus der Originaldatei: {blattname: [(bytes, anker-dict)]}."""
    out = defaultdict(list)
    with zipfile.ZipFile(pfad) as z:
        namen = set(z.namelist())

        def rels(teil):
            d, f = teil.rsplit("/", 1)
            rp = f"{d}/_rels/{f}.rels"
            if rp not in namen:
                return {}
            return {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read(rp)).findall("pr:Relationship", NS)}

        def ziel(basis, t):
            if t.startswith("/"):
                return t[1:]
            teile = basis.rsplit("/", 1)[0].split("/") + t.split("/")
            st = []
            for x in teile:
                if x == "..":
                    st.pop()
                elif x and x != ".":
                    st.append(x)
            return "/".join(st)

        wb = ET.fromstring(z.read("xl/workbook.xml"))
        wrels = rels("xl/workbook.xml")
        for sh in wb.find("m:sheets", NS):
            blatt = sh.get("name")
            sheet = ziel("xl/workbook.xml", wrels[sh.get(f"{{{NS['r']}}}id")])
            for rid, t in rels(sheet).items():
                if "drawings/drawing" not in t:
                    continue
                dpfad = ziel(sheet, t)
                drels = rels(dpfad)
                d = ET.fromstring(z.read(dpfad))
                for anc in list(d):
                    tag = anc.tag.split("}")[1]
                    blip = anc.find(".//a:blip", NS)
                    if blip is None or tag not in ("twoCellAnchor", "oneCellAnchor"):
                        continue
                    bild = z.read(ziel(dpfad, drels[blip.get(f"{{{NS['r']}}}embed")]))

                    def marker(el):
                        return {k: int(el.find(f"xdr:{k}", NS).text) for k in ("col", "colOff", "row", "rowOff")}
                    a = {"art": tag, "from": marker(anc.find("xdr:from", NS))}
                    if tag == "twoCellAnchor":
                        a["to"] = marker(anc.find("xdr:to", NS))
                    else:
                        ext = anc.find("xdr:ext", NS)
                        a["ext"] = (int(ext.get("cx")), int(ext.get("cy")))
                    out[blatt].append((bild, a))
    return out


def _bilder_setzen(ws, bilder):
    for daten, a in bilder:
        img = XlBild(io.BytesIO(daten))
        von = AnchorMarker(**{k: a["from"][k] for k in ("col", "colOff", "row", "rowOff")})
        if a["art"] == "twoCellAnchor":
            img.anchor = TwoCellAnchor(_from=von, to=AnchorMarker(**{k: a["to"][k] for k in ("col", "colOff", "row", "rowOff")}))
        else:
            img.anchor = OneCellAnchor(_from=von, ext=XDRPositiveSize2D(*a["ext"]))
        ws.add_image(img)


# ------------------------------------------------------------- Bauen
def bauen(con, quelle, ziel, n_zufall=20, n_schwer=10, saat=20261012):
    B = Bestand(con)
    quelle, ziel = Path(quelle), Path(ziel)
    ziel.mkdir(parents=True, exist_ok=True)
    gew, grundgesamtheit = auswahl(B, n_zufall, n_schwer, saat)
    farbe = dict(zip([f for f, _ in gew], farben(len(gew))))
    zeilen_von = {}                                          # (datei, blatt, zeile) -> [familien]
    familie_zeilen = {}
    for fid, _art in gew:
        z = [B.ort(e) for e in B.zeilen_der_familie(fid)]
        familie_zeilen[fid] = z
        for o in z:
            zeilen_von.setdefault(o, []).append(fid)
    ids_von = {B.ort(e): B.ids(e) for e in B.eintraege}
    belegt = {e for fid, _ in gew for e in B.zeilen_der_familie(fid)}
    kand_von = {}                                           # (datei, blatt, zeile) -> familie
    kand_meta = {}
    for fid, _art in gew:
        liste = []
        for eid, art_k, ziel_i in kandidaten(B, fid, belegt):
            belegt.add(eid)
            kand_von[B.ort(eid)] = fid
            liste.append({"ort": list(B.ort(eid)), "eintrag": eid, "art": art_k, "ziel": ziel_i})
        kand_meta[f"F{fid}"] = liste

    dateien = sorted({q["datei"] for q in B.quellen.values()})
    for datei in dateien:
        bilder = _bilder_lesen(quelle / datei)
        wb = openpyxl.load_workbook(quelle / datei)
        for ws in wb.worksheets:
            blaetter = {q["blatt"] for q in B.quellen.values() if q["datei"] == datei}
            if ws.title not in blaetter:
                continue
            letzte = ws.max_column
            kopf_fill = PatternFill("solid", fgColor="DDDDDD")
            for k, titel in enumerate(NEUE_SPALTEN, 1):
                c = ws.cell(1, letzte + k, titel)
                c.font = Font(bold=True); c.fill = kopf_fill; c.alignment = Alignment(wrap_text=True, vertical="top")
                ws.column_dimensions[openpyxl.utils.get_column_letter(letzte + k)].width = (12, 12, 14, 14, 40)[k - 1]
            dv = DataValidation(type="list", formula1='"richtig,falsch,unklar"', allow_blank=True)
            ws.add_data_validation(dv)
            for zeile in range(2, ws.max_row + 1):
                o = (datei, ws.title, zeile)
                pid, fid_txt = ids_von.get(o, ("", ""))
                ws.cell(zeile, letzte + 2, pid or None)
                ws.cell(zeile, letzte + 3, fid_txt or None)
                if o in zeilen_von:
                    fams = zeilen_von[o]
                    ws.cell(zeile, letzte + 1, ", ".join(f"F{f}" for f in fams))
                    fill = PatternFill("solid", fgColor=farbe[fams[0]])
                    for sp in range(1, letzte + len(NEUE_SPALTEN) + 1):
                        ws.cell(zeile, sp).fill = fill
                    dv.add(ws.cell(zeile, letzte + 4).coordinate)
                elif o in kand_von:                          # naheliegend, aber nicht zugeordnet: heller Ton, gleiche Frage
                    fid_k = kand_von[o]
                    ws.cell(zeile, letzte + 1, f"F{fid_k} (prüfen)")
                    fill = PatternFill("solid", fgColor=hell(farbe[fid_k]))
                    for sp in range(1, letzte + len(NEUE_SPALTEN) + 1):
                        ws.cell(zeile, sp).fill = fill
                    dv.add(ws.cell(zeile, letzte + 4).coordinate)
            ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(letzte + len(NEUE_SPALTEN))}{ws.max_row}"
            if not ws._images:                               # aeltere openpyxl verlieren Bilder; neuere lesen sie selbst
                _bilder_setzen(ws, bilder.get(ws.title, []))
        wb.save(ziel / datei)

    # Uebersicht
    ub = openpyxl.Workbook()
    ws = ub.active; ws.title = "Familien"
    ws.append(["Farbe", "Familie", "Mann", "Frau", "Trauung", "Zeilen in den Tabellen", "Fehlt eine Zeile?", "Welche (Tabelle, Zeile)?", "Kommentar", "Auswahl"])
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDDDDD"); c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"nein,ja"', allow_blank=True); ws.add_data_validation(dv)
    for fid, art in gew:
        f = B.fams[fid]
        nach = defaultdict(list)
        for d, b, z in familie_zeilen[fid]:
            nach[f"{d} · {b}"].append(z)
        zeilen = "\n".join(f"{k}: Zeile {', '.join(map(str, v))}" for k, v in nach.items())
        if kand_meta.get(f"F{fid}"):
            zeilen += "\nzum Prüfen (hell): " + "; ".join(f"{k['ort'][1]} Zeile {k['ort'][2]}" for k in kand_meta[f"F{fid}"])
        tr = f"{f['tr_tag'] or ''}.{f['tr_monat'] or ''}.{f['tr_jahr']}".strip(".") if f["tr_jahr"] else "keine (aus Taufen erschlossen)"
        ws.append(["", f"F{fid}", B.name(f["mann"]), B.name(f["frau"]), tr, zeilen, None, None, None, art])
        r = ws.max_row
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=farbe[fid])
        for c in ws[r]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        dv.add(ws.cell(r, 7).coordinate)
    for col, w in zip("ABCDEFGHIJ", (8, 9, 26, 26, 14, 60, 12, 28, 36, 18)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["J"].hidden = True
    ws.freeze_panes = "C2"
    an = ub.create_sheet("Anleitung")
    for t in ["In den Tabellen sind die Zeilen jeder Familie in einer Farbe markiert.",
              "Über den Filter-Pfeil in der Spalte „Stichprobe“ lassen sich nur diese Zeilen anzeigen.",
              "1. Gehört die farbige Zeile zur Familie? In der Spalte „Gehört dazu?“ richtig, falsch oder unklar wählen.",
              "   Hell gefärbte Zeilen hat das Programm NICHT zugeordnet, sie könnten aber dazugehören: dieselbe Frage.",
              "2. Fehlt eine Zeile, die zur Familie gehört? Hier im Blatt „Familien“ ja wählen und die Zeile nennen.",
              "Kommentare sind willkommen, Pflicht ist nichts."]:
        an.append([t])
    an.column_dimensions["A"].width = 110
    ub.save(ziel / "Stichprobe Übersicht.xlsx")
    (ziel / "stichprobe.json").write_text(json.dumps({"grundgesamtheit": grundgesamtheit, "familien": [[f, a] for f, a in gew],
        "zeilen": {f"F{f}": [list(o) for o in familie_zeilen[f]] for f, _ in gew},
        "eintraege": {f"F{f}": B.zeilen_der_familie(f) for f, _ in gew},
        "kandidaten": kand_meta, "mitglieder": {f"F{f}": {"mann": B.fams[f]["mann"], "frau": B.fams[f]["frau"]} for f, _ in gew},
        "nennungen": {str(e): {k: v["id"] for k, v in B.nennung[e].items()}
                      for e in {e for f, _ in gew for e in B.zeilen_der_familie(f)} | {k["eintrag"] for l in kand_meta.values() for k in l}},
        "erste_nennung": {str(i): B.erste_nennung.get(i) for i in {x for f, _ in gew for x in (B.fams[f]["mann"], B.fams[f]["frau"], *B.kinder.get(f, [])) if x}}}, ensure_ascii=False, indent=1), encoding="utf-8")
    return dict(familien=len(gew), zeilen=len(zeilen_von), kandidaten=len(kand_von), dateien=len(dateien), ziel=str(ziel),
                arten=Counter(a for _, a in gew), grundgesamtheit=grundgesamtheit)


# ------------------------------------------------------------- Auswerten
def _wilson(k, n, z=1.96):
    if n == 0:
        return (0.0, 0.0)
    p = k / n
    m = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    c = (p + z * z / (2 * n)) / (1 + z * z / n)
    return max(0.0, c - m), min(1.0, c + m)


def auswerten(ordner):
    """Ausgefuellte Dateien lesen -> dict mit Zaehlungen und Fehlerquoten, zufaellige und schwierige Familien getrennt."""
    ordner = Path(ordner)
    meta = json.loads((ordner / "stichprobe.json").read_text(encoding="utf-8"))
    art = {f"F{f}": a for f, a in meta["familien"]}
    urteil = {}                                              # (familie) -> Counter
    kand_urteil = {}                                         # familie -> Counter
    gefunden = defaultdict(list)                             # familie -> [(datei, blatt, zeile)] Kandidat "richtig"
    falsch_ort = defaultdict(list)                           # familie -> [(datei, blatt, zeile)]
    fremde = defaultdict(list)
    for datei in sorted(ordner.glob("*.xlsx")):
        if datei.name == "Stichprobe Übersicht.xlsx":
            continue
        wb = openpyxl.load_workbook(datei, data_only=True)
        for ws in wb.worksheets:
            kopf = {str(c.value).strip(): c.column for c in ws[1] if c.value}
            if "Stichprobe" not in kopf or "Gehört dazu?" not in kopf:
                continue
            for zeile in range(2, ws.max_row + 1):
                fam = ws.cell(zeile, kopf["Stichprobe"]).value
                if not fam:
                    continue
                u = (ws.cell(zeile, kopf["Gehört dazu?"]).value or "").strip().lower()
                if str(fam).endswith("(prüfen)"):
                    f = str(fam).split(" ")[0]
                    kand_urteil.setdefault(f, Counter())[u or "offen"] += 1
                    if u == "richtig":
                        gefunden[f].append((datei.name, ws.title, zeile))
                    continue
                for f in str(fam).split(", "):
                    if u == "falsch":
                        falsch_ort[f].append((datei.name, ws.title, zeile))
                    urteil.setdefault(f, Counter())[u or "offen"] += 1
                    if u == "falsch":
                        fremde[f].append(f"{datei.name} · {ws.title}: Zeile {zeile}")
    fehlt = {}
    ub = openpyxl.load_workbook(ordner / "Stichprobe Übersicht.xlsx", data_only=True)["Familien"]
    for r in ub.iter_rows(min_row=2, values_only=True):
        if r[1]:
            fehlt[r[1]] = ((r[6] or "").strip().lower() == "ja", r[7] or "")

    def gruppe(filter_):
        fs = [f for f in art if filter_(art[f])]
        richtig = sum(urteil.get(f, Counter())["richtig"] for f in fs)
        falsch = sum(urteil.get(f, Counter())["falsch"] for f in fs)
        unklar = sum(urteil.get(f, Counter())["unklar"] for f in fs)
        offen = sum(urteil.get(f, Counter())["offen"] for f in fs)
        geprueft = [f for f in fs if sum(urteil.get(f, Counter())[x] for x in ("richtig", "falsch")) > 0]
        fehlerhaft = [f for f in geprueft if urteil[f]["falsch"] or fehlt.get(f, (False,))[0]]
        mit_luecke = [f for f in geprueft if fehlt.get(f, (False,))[0] or kand_urteil.get(f, Counter())["richtig"]]
        fehlerhaft = [f for f in geprueft if urteil[f]["falsch"] or f in mit_luecke]
        k_richtig = sum(kand_urteil.get(f, Counter())["richtig"] for f in fs)
        k_falsch = sum(kand_urteil.get(f, Counter())["falsch"] for f in fs)
        k_offen = sum(kand_urteil.get(f, Counter())[x] for f in fs for x in ("unklar", "offen"))
        return dict(familien=len(fs), geprueft=len(geprueft), richtig=richtig, falsch=falsch, unklar=unklar, offen=offen,
                    zeilen_quote=falsch / (richtig + falsch) if richtig + falsch else None,
                    zeilen_intervall=_wilson(falsch, richtig + falsch),
                    familien_fehlerhaft=len(fehlerhaft), familien_quote=len(fehlerhaft) / len(geprueft) if geprueft else None,
                    familien_intervall=_wilson(len(fehlerhaft), len(geprueft)), mit_luecke=len(mit_luecke),
                    kand_richtig=k_richtig, kand_falsch=k_falsch, kand_offen=k_offen,
                    fehlend_quote=k_richtig / (k_richtig + richtig) if k_richtig + richtig else None,
                    fehlend_intervall=_wilson(k_richtig, k_richtig + richtig))
    return dict(zufall=gruppe(lambda a: a == "zufall"), schwierig=gruppe(lambda a: a != "zufall"),
                fremde=dict(fremde), fehlt={f: w for f, (j, w) in fehlt.items() if j},
                gefunden=dict(gefunden), falsch_ort=dict(falsch_ort), meta=meta)


def bericht(ordner):
    r = auswerten(ordner)
    meta = json.loads((Path(ordner) / "stichprobe.json").read_text(encoding="utf-8"))
    print(f"Grundgesamtheit: {meta.get('grundgesamtheit', '?')} Familien mit mindestens zwei Zeilen; die Quote gilt für sie.")
    pz = lambda x: "–" if x is None else f"{x * 100:.1f} %"
    for name, g in (("Zufällige Familien (Grundlage der Fehlerquote)", r["zufall"]), ("Schwierige Fälle (getrennt)", r["schwierig"])):
        print(f"\n{name}: {g['geprueft']} von {g['familien']} Familien geprüft")
        print(f"  Zeilen: richtig {g['richtig']}, falsch {g['falsch']}, unklar {g['unklar']}, ohne Urteil {g['offen']}")
        lo, hi = g["zeilen_intervall"]
        print(f"  Fremde Zeilen: {pz(g['zeilen_quote'])} (95 %: {pz(lo)} bis {pz(hi)})")
        lo, hi = g["familien_intervall"]
        print(f"  Familien mit Fehler (fremde oder fehlende Zeile): {g['familien_fehlerhaft']}, {pz(g['familien_quote'])} (95 %: {pz(lo)} bis {pz(hi)}); davon mit Lücke {g['mit_luecke']}")
        lo, hi = g["fehlend_intervall"]
        print(f"  Vorgelegte Kandidaten: gehört dazu {g['kand_richtig']}, gehört nicht dazu {g['kand_falsch']}, unklar/ohne Urteil {g['kand_offen']}")
        print(f"  Fehlende Zeilen: {pz(g['fehlend_quote'])} der zugehörigen Zeilen fehlten (95 %: {pz(lo)} bis {pz(hi)})")
    if r["fremde"]:
        print("\nFremde Zeilen:")
        for f, z in r["fremde"].items():
            print(f"  {f}: " + "; ".join(z))
    if r["fehlt"]:
        print("\nFehlende Zeilen:")
        for f, w in r["fehlt"].items():
            print(f"  {f}: {w}")
    return r


def uebernehmen(con, ordner):
    """Eindeutige Antworten als Entscheidungen speichern. -> (gespeichert, fuer die Pruefliste)."""
    from . import verknuepfen
    r = auswerten(ordner)
    meta = r["meta"]
    gespeichert, liste = [], []
    for f, orte in r["gefunden"].items():
        for o in orte:
            k = next((k for k in meta["kandidaten"].get(f, []) if tuple(k["ort"]) == tuple(o)), None)
            if not k:
                continue
            nenn = meta["nennungen"].get(str(k["eintrag"]), {})
            mitgl = meta["mitglieder"][f]
            if k["art"].startswith("Taufe"):
                for rolle, ident in (("vater", mitgl["mann"]), ("mutter", mitgl["frau"])):
                    ziel = meta["erste_nennung"].get(str(ident)) if ident else None
                    if nenn.get(rolle) and ziel:
                        verknuepfen.entscheiden_von_hand(con, nenn[rolle], "gleich", ziel)
                        gespeichert.append((f, o, rolle))
            else:
                ziel = meta["erste_nennung"].get(str(k["ziel"]))
                if nenn.get("verstorbener") and ziel:
                    verknuepfen.entscheiden_von_hand(con, nenn["verstorbener"], "gleich", ziel)
                    gespeichert.append((f, o, "verstorbener"))
    for f, orte in r["falsch_ort"].items():
        for o in orte:
            liste.append((f, o))
    # Begraebnis "falsch": der Verstorbene ist eine eigene Person
    eintrag_von_ort = {}
    for f, eids in meta["eintraege"].items():
        for e, o in zip(eids, meta["zeilen"][f]):
            eintrag_von_ort[(f, tuple(o))] = e
    rest = []
    for f, o in liste:
        e = eintrag_von_ort.get((f, tuple(o)))
        nenn = meta["nennungen"].get(str(e), {}) if e else {}
        if "verstorbener" in nenn:
            verknuepfen.entscheiden_von_hand(con, nenn["verstorbener"], "neu")
            gespeichert.append((f, o, "verstorbener neu"))
        else:
            rest.append((f, o))
    return gespeichert, rest
