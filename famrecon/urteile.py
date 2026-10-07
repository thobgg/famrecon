"""Pruefliste als Tabelle hin und zurueck: Urteile in Excel faellen, famrecon liest sie als Entscheidungen.

    famrecon urteile daten/projekt.db --schreiben pruefliste.xlsx [--alle]
    famrecon urteile daten/projekt.db --lesen pruefliste.xlsx          danach: famrecon verknuepfen

Eine Zeile je offenem Fall: Fall (Register, Jahr, Fundstelle, Rolle), die Nennung im Wortlaut, famrecons
Wahl mit Begruendung, die Kandidaten mit Nummer, dazu zwei leere Spalten URTEIL und BEMERKUNG.

    URTEIL   Nummer eines Kandidaten (die Zahl in eckigen Klammern) = die Nennung ist diese Person
             neu   = eigene Person, keiner der Kandidaten
             offen oder leer = keine Entscheidung
             zurueck = eine fruehere Entscheidung von Hand aufheben

Die Zuordnung laeuft ueber den stabilen Schluessel der Nennung (Datei|Blatt|Zeile|Rolle), wie in der
Oberflaeche; die Kandidatennummern gelten fuer den Stand, aus dem die Tabelle geschrieben wurde, deshalb
steht zu jedem Kandidaten auch der Schluessel einer seiner Nennungen in der Tabelle (Spalte KANDIDATEN).
So lassen sich Urteile auch nach einem Neubau einlesen.
"""
import json

import openpyxl

from . import verknuepfen

SPALTEN = ["FALL", "SCHLUESSEL", "REGISTER", "JAHR", "FUNDSTELLE", "ROLLE", "NENNUNG", "STUFE", "WAHL", "BEGRUENDUNG",
           "KANDIDATEN", "BISHER", "URTEIL", "BEMERKUNG"]


def _fund(con, eintrag_id):
    fl = {r["name"]: r["wert"] for r in con.execute("SELECT name, wert FROM feld WHERE eintrag=? AND wert IS NOT NULL", (eintrag_id,))}
    return fl.get("zitat") or " ".join(x for x in (fl.get("kb"), f"S. {fl['seite']}" if fl.get("seite") else None,
                                                    f"Nr. {fl['lfd_nr']}" if fl.get("lfd_nr") else None) if x) or ""


def _ident_text(con, ident_id):
    i = con.execute("SELECT name, vorname, geb_jahr, geb_praefix, tod_jahr, unbekannt FROM identitaet WHERE id=?", (ident_id,)).fetchone()
    if not i:
        return f"[{ident_id}]"
    return f"[{ident_id}] {i['name'] or ('NN' if i['unbekannt'] else '-')}, {i['vorname'] or '-'} *{i['geb_praefix'] or ''}{i['geb_jahr'] or '?'} †{i['tod_jahr'] or '?'}"


def _schluessel_von_ident(con, ident_id, nicht_person=None):
    """Stabiler Schluessel einer Nennung dieser Identitaet (nicht die fragliche selbst)."""
    r = con.execute("SELECT person FROM zuordnung WHERE ident=? AND person<>? ORDER BY person LIMIT 1", (ident_id, nicht_person or -1)).fetchone()
    return verknuepfen.person_schluessel(con, r["person"]) if r else None


def faelle(con, alle=False):
    """Die offenen Faelle wie in der Pruefliste: unsicher, neu mit Kandidaten, optional wahrscheinlich."""
    sql = ("SELECT z.*, p.pfad, p.roh, e.register, e.jahr, e.id eintrag_id, en.art hart, en.ziel hziel "
           "FROM zuordnung z JOIN person p ON p.id=z.person JOIN eintrag e ON e.id=p.eintrag JOIN quelle q ON q.id=e.quelle "
           "LEFT JOIN entscheidung en ON en.schluessel = q.datei||'|'||q.blatt||'|'||e.zeile||'|'||p.pfad "
           "WHERE (z.stufe IN ('unsicher'" + (",'wahrscheinlich'" if alle else "") + ") OR (z.stufe='neu' AND z.alternativen<>'[]') OR en.art IS NOT NULL) "
           "ORDER BY e.jahr, e.id, p.id")
    return con.execute(sql).fetchall()


def tabelle_schreiben(con, pfad, alle=False):
    """Pruefliste als xlsx. Gibt die Zahl der Faelle zurueck."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pruefliste"
    ws.append(SPALTEN)
    n = 0
    for r in faelle(con, alle):
        kandidaten = [(r["ident"], r["punkte"], r["grund"])] + [tuple(a) for a in json.loads(r["alternativen"] or "[]")]
        if r["stufe"] == "neu":                            # die eigene neue Person ist kein Kandidat
            kandidaten = kandidaten[1:]
        k_text = "\n".join(f"{_ident_text(con, k[0])} ({k[1]} P.: {k[2]})" for k in kandidaten)
        k_keys = ";".join(f"{k[0]}={_schluessel_von_ident(con, k[0], r['person']) or ''}" for k in kandidaten)
        wahl = _ident_text(con, r["ident"]) if r["stufe"] != "neu" else "eigene Person (unentschieden)"
        bisher = (r["hart"] + (" " + r["hziel"] if r["hziel"] else "")) if r["hart"] else ""
        fall = f"{r['register']} {r['jahr'] or '?'} · {_fund(con, r['eintrag_id'])} · {r['pfad']}"
        ws.append([fall, verknuepfen.person_schluessel(con, r["person"]), r["register"], r["jahr"], _fund(con, r["eintrag_id"]), r["pfad"],
                   r["roh"] or "", r["stufe"], wahl, r["grund"] or "", k_text + "\n" + "#" + k_keys, bisher, "", ""])
        n += 1
    for col, breite in zip("ABCDEFGHIJKLMN", (34, 10, 8, 6, 26, 18, 40, 12, 40, 36, 60, 14, 10, 30)):
        ws.column_dimensions[col].width = breite
    ws.column_dimensions["B"].hidden = True
    ws.freeze_panes = "A2"
    wb.save(pfad)
    return n


def tabelle_lesen(con, pfad):
    """Urteile aus der Tabelle als Entscheidungen speichern. -> dict(gleich=, neu=, zurueck=, uebergangen=, fehler=[])"""
    ws = openpyxl.load_workbook(pfad, read_only=True, data_only=True).active
    zeilen = ws.iter_rows(values_only=True)
    kopf = [str(c).strip().upper() if c else "" for c in next(zeilen)]
    try:
        i_key, i_urteil, i_kand = kopf.index("SCHLUESSEL"), kopf.index("URTEIL"), kopf.index("KANDIDATEN")
    except ValueError:
        return dict(gleich=0, neu=0, zurueck=0, uebergangen=0, fehler=["Tabelle ohne die Spalten SCHLUESSEL, URTEIL, KANDIDATEN"])
    person_von = {verknuepfen.person_schluessel(con, r["id"]): r["id"] for r in con.execute("SELECT id FROM person")}
    st = dict(gleich=0, neu=0, zurueck=0, uebergangen=0, fehler=[])
    for zeile in zeilen:
        if not zeile or i_key >= len(zeile):
            continue
        key, urteil = zeile[i_key], (str(zeile[i_urteil]).strip() if i_urteil < len(zeile) and zeile[i_urteil] is not None else "")
        if not key or not urteil or urteil.lower() in ("offen", "?", "-"):
            st["uebergangen"] += 1
            continue
        person = person_von.get(str(key))
        if person is None:
            st["fehler"].append(f"{key}: Nennung nicht mehr im Projekt"); continue
        u = urteil.lower()
        if u in ("neu", "eigene person", "keiner"):
            verknuepfen.entscheiden_von_hand(con, person, "neu"); st["neu"] += 1
        elif u in ("zurueck", "zurück", "loeschen", "löschen"):
            verknuepfen.entscheiden_von_hand(con, person, None); st["zurueck"] += 1
        else:
            nummer = u.strip("[] ")
            kand = str(zeile[i_kand] or "")
            keys = dict(x.split("=", 1) for x in kand.rsplit("#", 1)[-1].split(";") if "=" in x) if "#" in kand else {}
            ziel_key = keys.get(nummer)
            ziel = person_von.get(ziel_key) if ziel_key else None
            if ziel is None:                              # Kandidat ohne Schluessel: ueber die aktuelle Identitaetsnummer
                r = con.execute("SELECT person FROM zuordnung WHERE ident=? AND person<>? ORDER BY person LIMIT 1", (nummer, person)).fetchone() if nummer.isdigit() else None
                ziel = r["person"] if r else None
            if ziel is None:
                st["fehler"].append(f"{key}: Urteil {urteil!r} passt zu keinem Kandidaten"); continue
            verknuepfen.entscheiden_von_hand(con, person, "gleich", ziel); st["gleich"] += 1
    return st
