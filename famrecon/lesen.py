"""Tabellen einlesen: Excel-Zeilen nach Spaltenzuordnung in eintrag/feld.

    famrecon spalten datei.xlsx                 Blaetter und Ueberschriften zeigen
    famrecon bau zuordnung.toml datei.xlsx ...  in projekt.db einlesen

Die Zuordnung (TOML) nennt je Register das Blatt und je Spalte den Feldnamen.
Mehrere Spalten koennen ein Feld bilden ("S-AN" = "paten"); die Werte werden
mit " | " verbunden. Was in `leer` steht ("k.A.", "—"), gilt als nicht
angegeben: `wert` bleibt leer, `roh` behaelt den Zellinhalt.
"""
import datetime as dt
import re
import tomllib
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from . import normalform
from .i18n import _

TRENNER = " | "


def zuordnung_laden(pfad):
    """TOML-Zuordnung lesen; Spaltenbereiche wie 'S-AN' werden zu Indexlisten aufgeloest, `leer` bekommt einen Standard."""
    with open(pfad, "rb") as f:
        z = tomllib.load(f)
    z.setdefault("allgemein", {}).setdefault("leer", [])
    for name, reg in z.get("register", {}).items():
        reg["_spalten"] = _spalten_aufloesen(reg["spalten"])
        reg.setdefault("register", name)          # Abschnitt taufe_2 -> Register taufe
    return z


def _spalten_aufloesen(spalten):
    """{"A": "x", "S-AN": "paten"} -> [(feldname, [spaltenindex, ...]), ...]"""
    out = []
    for schluessel, feld in spalten.items():
        if "-" in schluessel:
            von, bis = schluessel.split("-", 1)
            idx = list(range(column_index_from_string(von.strip()),
                             column_index_from_string(bis.strip()) + 1))
        else:
            idx = [column_index_from_string(schluessel.strip())]
        out.append((feld, idx))
    return out


def bereinigen(zelle, leerwoerter):
    """Zellinhalt -> (wert, roh). Datumszellen werden ISO, Zahlen Text."""
    if zelle is None:
        return None, None
    if isinstance(zelle, dt.datetime):
        roh = zelle.isoformat()
        return zelle.date().isoformat(), roh
    if isinstance(zelle, dt.date):
        return zelle.isoformat(), zelle.isoformat()
    if isinstance(zelle, float) and zelle.is_integer():
        zelle = int(zelle)
    roh = str(zelle)
    wert = re.sub(r"\s+", " ", roh).strip()
    if not wert or wert in leerwoerter:
        return None, roh
    return wert, roh


def jahr_aus(wert):
    """Jahr (vier Ziffern am Anfang) aus einem Datumswert oder None."""
    m = re.match(r"(\d{4})", wert or "")
    return int(m.group(1)) if m else None


def _csv_lesen(pfad):
    """CSV/TSV: Trennzeichen und Kodierung raten. -> [tuple, ...] inklusive Kopfzeile."""
    import csv
    roh = Path(pfad).read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = roh.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialekt = csv.Sniffer().sniff(text[:4000], delimiters=";,\t|")
    except csv.Error:
        dialekt = csv.excel
        dialekt.delimiter = ";" if text.count(";") > text.count(",") else ","
    return [tuple((c if c != "" else None) for c in r) for r in csv.reader(text.splitlines(), dialekt)]


_CACHE = {}


def _tabellen(pfad):
    """-> {blattname: [zeilen]} fuer xlsx, oder ein Blatt (Dateiname ohne Endung) fuer csv/tsv.
    Einmal gelesen, dann aus dem Speicher (Schluessel: Pfad und Aenderungszeit), weil Zuordnung
    und Einlesen dieselbe Datei mehrfach brauchen."""
    pfad = Path(pfad)
    schluessel = (str(pfad.resolve()), pfad.stat().st_mtime_ns)
    if schluessel in _CACHE:
        return _CACHE[schluessel]
    if pfad.suffix.lower() in (".csv", ".tsv", ".txt"):
        out = {pfad.stem: _csv_lesen(pfad)}
    else:
        wb = openpyxl.load_workbook(pfad, read_only=True, data_only=True)
        out = {ws.title: list(ws.iter_rows(values_only=True)) for ws in wb.worksheets}
        wb.close()
    _CACHE.clear()                      # nur die zuletzt gelesene Datei behalten: Speicher
    _CACHE[schluessel] = out
    return out


def blaetter(pfad):
    """Blaetter mit Kopfzeile und Zeilenzahl, fuer `famrecon spalten`."""
    out = []
    for titel, zeilen in _tabellen(pfad).items():
        kopf = zeilen[0] if zeilen else ()
        n = sum(1 for r in zeilen[1:] if any(c is not None for c in r))
        out.append((titel, n, [(get_column_letter(i + 1), re.sub(r"\s+", " ", str(k)).strip())
                               for i, k in enumerate(kopf) if k is not None]))
    return out


def zeilen_lesen(pfad, blatt):
    """Alle nicht leeren Datenzeilen eines Blatts: (zeilennummer, tuple)."""
    tabellen = _tabellen(pfad)
    if blatt not in tabellen:
        raise SystemExit(_("Blatt '{blatt}' fehlt in {datei}; vorhanden: {blaetter}").format(blatt=blatt, datei=pfad, blaetter=list(tabellen)))
    for nr, r in enumerate(tabellen[blatt][1:], start=2):
        if any(c is not None for c in r):
            yield nr, r


def zeile_zu_feldern(reihe, spalten, leerwoerter):
    """Eine Tabellenzeile -> {feldname: (wert, roh)} nach der Zuordnung."""
    felder = {}
    for feld, idx in spalten:
        werte, rohs = [], []
        for i in idx:
            zelle = reihe[i - 1] if i - 1 < len(reihe) else None
            w, r = bereinigen(zelle, leerwoerter)
            if w:
                werte.append(w)
            if r is not None:
                rohs.append(r)
        if werte or rohs:
            alt = felder.get(feld)
            if alt:
                werte = ([alt[0]] if alt[0] else []) + werte
                rohs = ([alt[1]] if alt[1] else []) + rohs
            felder[feld] = (TRENNER.join(werte) if werte else None,
                            TRENNER.join(rohs) if rohs else None)
    return felder


def unbekannte_felder(zuordnung):
    """Feldnamen der Zuordnung, die der Katalog nicht kennt: [(register, feld)]."""
    from . import katalog
    out = []
    for register, reg in zuordnung.get("register", {}).items():
        out += [(register, f) for f, _ in reg["_spalten"]
                if register not in katalog.REGISTER or not katalog.bekannt(register, f)]
    return out


def einlesen(con, zuordnung, dateien):
    """Alle Register aus den Dateien in die DB. Gibt {register: zeilen} zurueck."""
    leer = set(zuordnung["allgemein"]["leer"])
    zaehler = {}
    jetzt = dt.datetime.now().isoformat(timespec="seconds")
    for tabelle in ("zuordnung", "kind", "familie", "identitaet", "person"):
        con.execute(f"DELETE FROM {tabelle}")
    for datei in dateien:
        datei = Path(datei)
        wb_blaetter = [b for b, _, _ in blaetter(datei)]
        for _abschnitt, reg in zuordnung["register"].items():
            register = reg["register"]
            if reg["blatt"] not in wb_blaetter:
                continue
            if reg.get("datei") and reg["datei"] != datei.name:
                continue
            con.execute("DELETE FROM quelle WHERE datei=? AND blatt=?",
                        (datei.name, reg["blatt"]))
            cur = con.execute(
                "INSERT INTO quelle(datei, blatt, register, zeilen, gelesen_am) "
                "VALUES (?,?,?,0,?)", (datei.name, reg["blatt"], register, jetzt))
            qid = cur.lastrowid
            n = 0
            for zeilennr, reihe in zeilen_lesen(datei, reg["blatt"]):
                felder = zeile_zu_feldern(reihe, reg["_spalten"], leer)
                jmt = None
                for leit in reg.get("leitdatum", []):
                    jmt = normalform.datum_zerlegen((felder.get(leit) or (None,))[0])
                    if jmt:
                        break
                if not jmt and felder.get("jahr", (None,))[0]:
                    jmt = normalform.datum_zerlegen(felder["jahr"][0])
                jahr, monat, tag = jmt or (None, None, None)
                eid = con.execute(
                    "INSERT INTO eintrag(quelle, register, zeile, jahr, monat, tag) VALUES (?,?,?,?,?,?)",
                    (qid, register, zeilennr, jahr, monat, tag)).lastrowid
                con.executemany(
                    "INSERT INTO feld(eintrag, name, wert, roh) VALUES (?,?,?,?)",
                    [(eid, name, w, r) for name, (w, r) in felder.items()])
                n += 1
            con.execute("UPDATE quelle SET zeilen=? WHERE id=?", (n, qid))
            zaehler[register] = zaehler.get(register, 0) + n
    con.execute("INSERT OR REPLACE INTO einstellung(name, wert) VALUES ('kennungen', ?)",
                ("1" if zuordnung["allgemein"].get("kennungen") else "0"))
    con.commit()
    return zaehler
